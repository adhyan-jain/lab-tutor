"""Usage metrics are tracked for every role; student aggregates stay student-only.

Staff test traffic (faculty, co-faculty, admin) is tagged FACULTY_TEST /
ADMIN_TEST on purpose so it never mixes into student research numbers. The
Activity view can list it separately, labelled, and the export has its own
"Staff activity" sheet.
"""

from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook

from backend.auth import session as session_cookie

pytestmark = pytest.mark.asyncio

ADMIN_EMAIL = "adhyanjain2006@gmail.com"


def auth(token: str) -> dict[str, str]:
    return {"Cookie": f"{session_cookie.COOKIE_NAME}={token}"}


@pytest.fixture
def llm(fake_llm, monkeypatch):
    fake_llm.reply = "HOMO is the highest occupied molecular orbital."
    monkeypatch.setattr("backend.retrieval.pipeline.get_backend", lambda: fake_llm)
    return fake_llm


async def _chat(client, token, cid, text="What is HOMO?"):
    resp = await client.post(
        "/api/chat/messages",
        json={"classroom_id": cid, "experiment_id": "exp07", "message": text},
        headers=auth(token),
    )
    assert resp.status_code == 200, resp.text


async def _world(client, make_user):
    """One class, a student, the professor, a promoted co-faculty and an admin,
    each having sent exactly one prompt."""
    prof_user, prof = await make_user("owner.prof@vit.ac.in", "Prof Owner")
    created = await client.post("/api/classrooms", json={"name": "Sec A"}, headers=auth(prof))
    cid, code = created.json()["id"], created.json()["student_join_code"]
    await client.post(
        f"/api/classrooms/{cid}/sessions/start", json={"experiment_id": "exp07"}, headers=auth(prof)
    )
    _, student = await make_user("real.student@vitstudent.ac.in", "Real Student")
    await client.post("/api/classrooms/join", json={"join_code": code}, headers=auth(student))
    co_user, co_token = await make_user("co.faculty@vitstudent.ac.in", "Co Faculty")
    await client.post("/api/classrooms/join", json={"join_code": code}, headers=auth(co_token))
    await client.post(
        f"/api/classrooms/{cid}/promote", json={"student_user_id": co_user.id}, headers=auth(prof)
    )
    _, admin = await make_user(ADMIN_EMAIL, "Admin")
    for token in (student, prof, co_token, admin):
        await _chat(client, token, cid)
    return prof, admin, cid


async def _activity(client, token, cid, role=None):
    url = f"/api/dashboard/classrooms/{cid}/activity" + (f"?role={role}" if role else "")
    resp = await client.get(url, headers=auth(token))
    assert resp.status_code == 200, resp.text
    return resp.json()


def _by_email(payload):
    return {row["email"]: row for row in payload["students"]}


async def test_default_view_is_students_only(client, make_user, llm):
    prof, _, cid = await _world(client, make_user)
    payload = await _activity(client, prof, cid)
    assert list(_by_email(payload)) == ["real.student@vitstudent.ac.in"]
    assert payload["totals"]["prompts"] == 1
    assert payload["role_filter"] == "student"


async def test_everyone_view_lists_staff_with_their_role_and_prompts(client, make_user, llm):
    prof, _, cid = await _world(client, make_user)
    payload = await _activity(client, prof, cid, "all")
    rows = _by_email(payload)
    assert rows["real.student@vitstudent.ac.in"]["role"] == "student"
    assert rows["owner.prof@vit.ac.in"]["role"] == "faculty"
    assert rows["co.faculty@vitstudent.ac.in"]["role"] == "co-faculty"
    assert rows[ADMIN_EMAIL]["role"] == "admin"
    assert all(row["prompts_total"] == 1 for row in rows.values())
    assert payload["totals"]["prompts"] == 4
    assert payload["by_role"] == {
        "student": {"users": 1, "prompts": 1},
        "faculty": {"users": 1, "prompts": 1},
        "co-faculty": {"users": 1, "prompts": 1},
        "admin": {"users": 1, "prompts": 1},
    }


async def test_role_filters_pick_the_right_people(client, make_user, llm):
    prof, _, cid = await _world(client, make_user)
    faculty = _by_email(await _activity(client, prof, cid, "faculty"))
    assert set(faculty) == {"owner.prof@vit.ac.in", "co.faculty@vitstudent.ac.in"}
    admin = _by_email(await _activity(client, prof, cid, "admin"))
    assert set(admin) == {ADMIN_EMAIL}


async def test_staff_rows_carry_the_same_model_metrics(client, make_user, llm):
    prof, _, cid = await _world(client, make_user)
    row = _by_email(await _activity(client, prof, cid, "all"))["owner.prof@vit.ac.in"]
    for key in ("prompt_tokens", "completion_tokens", "thinking_tokens", "cached_tokens", "llm_calls",
                "retries", "fallback_replies", "logins", "active_seconds", "experiments"):
        assert key in row
    assert row["experiments"] == ["exp07"]


async def test_a_bad_role_value_is_rejected(client, make_user, llm):
    prof, _, cid = await _world(client, make_user)
    resp = await client.get(f"/api/dashboard/classrooms/{cid}/activity?role=everything", headers=auth(prof))
    assert resp.status_code == 422


async def test_students_cannot_read_the_activity_view_at_all(client, make_user, llm):
    _, _, cid = await _world(client, make_user)
    _, outsider = await make_user("outsider@vitstudent.ac.in")
    for role in (None, "all"):
        url = f"/api/dashboard/classrooms/{cid}/activity" + (f"?role={role}" if role else "")
        assert (await client.get(url, headers=auth(outsider))).status_code in (403, 404)


async def test_export_has_a_staff_sheet_and_keeps_staff_out_of_student_sheets(client, make_user, llm):
    prof, _, cid = await _world(client, make_user)
    resp = await client.get(f"/api/dashboard/classrooms/{cid}/research-export.xlsx", headers=auth(prof))
    assert resp.status_code == 200
    book = load_workbook(io.BytesIO(resp.content))
    assert "Staff activity" in book.sheetnames

    staff = {row[1]: row for row in list(book["Staff activity"].iter_rows(values_only=True))[1:]}
    assert set(staff) == {"owner.prof@vit.ac.in", "co.faculty@vitstudent.ac.in", ADMIN_EMAIL}
    assert staff[ADMIN_EMAIL][2] == "admin" and staff[ADMIN_EMAIL][5] == 1

    prompt_emails = {row[1] for row in list(book["Prompt counts"].iter_rows(values_only=True))[1:]}
    assert prompt_emails == {"real.student@vitstudent.ac.in"}
    roster_emails = {row[1] for row in list(book["Roster"].iter_rows(values_only=True))[1:]}
    assert roster_emails == {"real.student@vitstudent.ac.in"}
