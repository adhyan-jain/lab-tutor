"""Research-data export: login sessions, prompt counts, trajectory, marks
combined into one workbook for faculty/admin. See CLAUDE.md's testing
philosophy -- expected values below are computed by hand, not invented.
"""

from __future__ import annotations

import datetime as dt
import io

import pytest
from openpyxl import load_workbook

from backend.auth import session as session_cookie
from backend.models import ActorType, ChatMessage, ChatMessageKind, LoginSession

pytestmark = pytest.mark.asyncio


def auth(token: str) -> dict[str, str]:
    return {"Cookie": f"{session_cookie.COOKIE_NAME}={token}"}


async def _make_classroom_with_student(client, make_user, faculty_email="prof.export@vit.ac.in"):
    _, faculty_token = await make_user(faculty_email, "Prof Export")
    created = await client.post(
        "/api/classrooms", json={"name": "Export Test Section"}, headers=auth(faculty_token)
    )
    assert created.status_code == 201
    classroom_id = created.json()["id"]
    student_code = created.json()["student_join_code"]

    student, student_token = await make_user("student.export@vitstudent.ac.in", "Student Export")
    joined = await client.post(
        "/api/classrooms/join", json={"join_code": student_code}, headers=auth(student_token)
    )
    assert joined.status_code == 200
    return classroom_id, faculty_token, student, student_token


async def test_export_includes_sessions_and_prompt_counts(client, make_user, db):
    classroom_id, faculty_token, student, _ = await _make_classroom_with_student(
        client, make_user
    )

    login_at = dt.datetime.now(dt.timezone.utc) - dt.timedelta(hours=1)
    logout_at = dt.datetime.now(dt.timezone.utc)
    db.add(
        LoginSession(
            user_id=student.id,
            login_at=login_at,
            last_seen_at=logout_at,
            logout_at=logout_at,
            end_reason="logout",
        )
    )
    for _ in range(3):
        db.add(
            ChatMessage(
                student_id=student.id,
                classroom_id=classroom_id,
                experiment_id="exp07",
                kind=ChatMessageKind.QA,
                actor_type=ActorType.STUDENT,
                author="student",
                content="how do I build methane in gabedit",
            )
        )
    await db.commit()

    resp = await client.get(
        f"/api/dashboard/classrooms/{classroom_id}/research-export.xlsx",
        headers=auth(faculty_token),
    )
    assert resp.status_code == 200
    workbook = load_workbook(io.BytesIO(resp.content))
    assert workbook.sheetnames == ["Roster", "Sessions", "Prompt counts", "Trajectory", "Marks"]

    sessions_sheet = workbook["Sessions"]
    session_row = [cell.value for cell in sessions_sheet[2]]
    assert session_row[1] == "student.export@vitstudent.ac.in"
    assert session_row[3] is not None  # logout at, explicit
    assert session_row[5] == "logout"

    prompts_sheet = workbook["Prompt counts"]
    prompt_row = [cell.value for cell in prompts_sheet[2]]
    assert prompt_row[1] == "student.export@vitstudent.ac.in"
    assert prompt_row[2] == "exp07"
    assert prompt_row[4] == "qa"
    assert prompt_row[5] == 3


async def test_export_marks_sheet_matches_entered_values(client, make_user, db):
    classroom_id, faculty_token, student, _ = await _make_classroom_with_student(
        client, make_user
    )
    await client.post(
        f"/api/marks/classrooms/{classroom_id}/experiments/exp07",
        json={
            "entries": [
                {
                    "student_id": student.id,
                    "pre_test_marks": 6, "pre_test_max": 20,
                    "post_test_marks": 15, "post_test_max": 20,
                }
            ]
        },
        headers=auth(faculty_token),
    )

    resp = await client.get(
        f"/api/dashboard/classrooms/{classroom_id}/research-export.xlsx",
        headers=auth(faculty_token),
    )
    assert resp.status_code == 200
    workbook = load_workbook(io.BytesIO(resp.content))
    marks_row = [cell.value for cell in workbook["Marks"][2]]
    assert marks_row[0] == "exp07"
    assert marks_row[3] == 6
    assert marks_row[5] == 15
    assert marks_row[7] == 9  # gain


async def test_export_forbidden_for_a_non_member_faculty(client, make_user):
    classroom_id, _faculty_token, _student, _ = await _make_classroom_with_student(
        client, make_user
    )
    _, outsider_token = await make_user("prof.outsider@vit.ac.in", "Outsider Prof")
    resp = await client.get(
        f"/api/dashboard/classrooms/{classroom_id}/research-export.xlsx",
        headers=auth(outsider_token),
    )
    assert resp.status_code == 404
